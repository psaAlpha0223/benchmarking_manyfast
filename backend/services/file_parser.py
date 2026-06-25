import base64
import io
from dataclasses import dataclass
from typing import Optional

import openpyxl
import pdfplumber
from docx import Document

from utils.text_cleaner import clean_text

IMAGE_MEDIA_TYPES = {
    "png": "image/png",
    "jpg": "image/jpeg",
    "jpeg": "image/jpeg",
}


@dataclass
class ParsedFile:
    filename: str
    kind: str  # "text" | "image" | "unsupported"
    text: Optional[str] = None
    base64_data: Optional[str] = None
    media_type: Optional[str] = None


def _extract_pdf_text(content: bytes) -> str:
    chunks = []
    with pdfplumber.open(io.BytesIO(content)) as pdf:
        for page in pdf.pages:
            text = page.extract_text()
            if text:
                chunks.append(text)
    return clean_text("\n".join(chunks))


def _extract_docx_text(content: bytes) -> str:
    document = Document(io.BytesIO(content))
    paragraphs = [p.text for p in document.paragraphs if p.text]
    return clean_text("\n".join(paragraphs))


def _extract_xlsx_text(content: bytes) -> str:
    workbook = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    chunks = []
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                chunks.append(" | ".join(cells))
    return clean_text("\n".join(chunks))


def parse_file(filename: str, content: bytes) -> ParsedFile:
    ext = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""

    if ext == "pdf":
        return ParsedFile(filename=filename, kind="text", text=_extract_pdf_text(content))

    if ext == "docx":
        return ParsedFile(filename=filename, kind="text", text=_extract_docx_text(content))

    if ext == "xlsx":
        return ParsedFile(filename=filename, kind="text", text=_extract_xlsx_text(content))

    if ext in IMAGE_MEDIA_TYPES:
        return ParsedFile(
            filename=filename,
            kind="image",
            base64_data=base64.b64encode(content).decode("ascii"),
            media_type=IMAGE_MEDIA_TYPES[ext],
        )

    # .hwp, .xls: 현재 기술 스택(requirements.txt)에 파싱 라이브러리가 없어
    # 텍스트 추출 없이 파일명만 Claude에게 전달한다.
    return ParsedFile(filename=filename, kind="unsupported")


def gather_file_context(file_paths: Optional[list[str]]) -> tuple[str, list[dict]]:
    from services import supabase_service

    text_parts = []
    image_blocks = []

    for path in file_paths or []:
        content = supabase_service.download_file(path)
        filename = path.rsplit("/", 1)[-1]
        parsed = parse_file(filename, content)

        if parsed.kind == "text":
            text_parts.append(f"[첨부파일: {parsed.filename}]\n{parsed.text}")
        elif parsed.kind == "image":
            image_blocks.append(
                {"media_type": parsed.media_type, "base64_data": parsed.base64_data}
            )
        else:
            text_parts.append(f"[첨부파일: {parsed.filename}] (텍스트 추출 미지원 형식)")

    return "\n\n".join(text_parts), image_blocks
